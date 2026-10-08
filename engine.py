from contextlib import contextmanager
import csv, io, json, hashlib, sqlite3, os, base64
from decimal import Decimal, InvalidOperation
DB=os.getenv('DATA_DB','data.sqlite3')
SAMPLE='Code,Product,Quantity,Price\nSKU-101,Canvas tote,24,18.50\nSKU-102,Travel mug,12,24.90\nSKU-103,Notebook,40,8.00\nSKU-102,Duplicate mug,9,24.90\nSKU-104,Desk lamp,-2,39.00'
@contextmanager
def connect():
    c=sqlite3.connect(DB,timeout=10);c.row_factory=sqlite3.Row
    c.executescript('CREATE TABLE IF NOT EXISTS products(sku TEXT PRIMARY KEY,name TEXT NOT NULL,quantity INTEGER NOT NULL,price TEXT NOT NULL); CREATE TABLE IF NOT EXISTS imports(digest TEXT PRIMARY KEY,count INTEGER NOT NULL);')
    try:
        yield c
        c.commit()
    except Exception:
        c.rollback();raise
    finally:
        c.close()
def state():
    with connect() as c: return {'sample':SAMPLE,'products':[dict(r) for r in c.execute('SELECT * FROM products ORDER BY sku')], 'imports':c.execute('SELECT COUNT(*) FROM imports').fetchone()[0]}
def rows(data):
    if data.get('format','csv')=='xlsx':
        from openpyxl import load_workbook
        raw=base64.b64decode(data['file'],validate=True)
        import zipfile
        with zipfile.ZipFile(io.BytesIO(raw)) as z:
            if sum(i.file_size for i in z.infolist())>20000000: raise ValueError('Expanded XLSX exceeds 20 MB')
        wb=load_workbook(io.BytesIO(raw),read_only=True,data_only=True)
        try:
            iterator=wb.active.iter_rows(values_only=True);headers=[str(x or '').strip() for x in next(iterator)]
            result=[]
            for row in iterator:
                if len(result)>=1000: raise ValueError('Maximum 1000 rows')
                result.append(dict(zip(headers,row)))
            return result
        finally: wb.close()
    reader=csv.DictReader(io.StringIO(data['csv']))
    if not reader.fieldnames or len(set(reader.fieldnames))!=len(reader.fieldnames): raise ValueError('Missing or duplicate column names')
    result=[]
    for row in reader:
        if len(result)>=1000: raise ValueError('Maximum 1000 rows')
        result.append(row)
    return result
def preview(data, existing=None):
    mapping=data.get('mapping',{'sku':'Code','name':'Product','quantity':'Quantity','price':'Price'})
    required=('sku','name','quantity','price')
    if set(mapping)!=set(required): raise ValueError('Map all four fields')
    if len(set(mapping.values()))!=4: raise ValueError('Each field needs a distinct source column')
    if existing is None:
        with connect() as c: existing={r[0] for r in c.execute('SELECT sku FROM products')}
    seen=set();result=[]
    for n,row in enumerate(rows(data),2):
        errors=[];item={k:str(row.get(mapping[k],'') if row.get(mapping[k]) is not None else '').strip() for k in required}
        if not item['sku'] or len(item['sku'])>64: errors.append('SKU required (max 64 characters)')
        if not item['name'] or len(item['name'])>200: errors.append('Product name required (max 200 characters)')
        try:
            item['quantity']=int(item['quantity'])
            if not 0<=item['quantity']<=1000000: raise ValueError()
        except ValueError: errors.append('Quantity must be an integer from 0 to 1000000')
        try:
            price=Decimal(item['price'])
            if not price.is_finite() or not 0<=price<=1000000 or price.as_tuple().exponent < -2: raise ValueError()
            item['price']=str(price.quantize(Decimal('.01')))
        except (InvalidOperation,ValueError): errors.append('Price must be nonnegative with at most 2 decimals')
        if item['sku'] in seen or item['sku'] in existing: errors.append('Duplicate SKU')
        seen.add(item['sku']);result.append({'line':n,'item':item,'errors':errors})
    return {'rows':result,'valid':sum(not r['errors'] for r in result),'invalid':sum(bool(r['errors']) for r in result)}
def action(path,data):
    if path=='/api/preview': return preview(data)
    if path!='/api/import': raise ValueError('Unknown operation')
    # Digest includes mapping: retries never create another batch.
    digest=hashlib.sha256(json.dumps(data,sort_keys=True).encode()).hexdigest()
    with connect() as c:
        c.execute('BEGIN IMMEDIATE')
        prior=c.execute('SELECT count FROM imports WHERE digest=?',(digest,)).fetchone()
        if prior: return {'imported':prior[0],'replayed':True}
        report=preview(data,{r[0] for r in c.execute('SELECT sku FROM products')});valid=[r['item'] for r in report['rows'] if not r['errors']]
        if not valid: raise ValueError('No valid new rows to import')
        c.executemany('INSERT INTO products VALUES(:sku,:name,:quantity,:price)',valid)
        c.execute('INSERT INTO imports VALUES(?,?)',(digest,len(valid)))
        return {'imported':len(valid),'skipped':report['invalid'],'replayed':False}
